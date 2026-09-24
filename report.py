"""
report.py — the Markdown / Word / PDF report and the console summary.

Everything reads through db.load_listings(), so the report shows exactly what
the dashboard and the alerts show: no expired, duplicate, stale or filtered
listings, one score per listing.
"""
from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from datetime import datetime

from common import COUNTRY_NAMES, COUNTRY_ORDER, LOG, days_left, utcnow
from db import hidden_category, load_listings, source_health

HERE = os.path.dirname(os.path.abspath(__file__))

SECTION_NAMES = {
    "imoveis": "Imóveis / Real Estate",
    "ouro_joias": "Ouro & Joias (Gold & Jewelry)",
    "outros": "Outros (Vehicles, Equipment, etc.)",
}


def _countries(items) -> list[str]:
    present = {it.get("country") or "PT" for it in items}
    return [c for c in COUNTRY_ORDER if c in present] + sorted(present - set(COUNTRY_ORDER))


def _loc(item) -> str:
    return ", ".join(filter(None, [item.get("concelho"), item.get("district")]))


def _money(v) -> str:
    return f"€{v:,.0f}" if v else "?"


def md_cell(text, limit: int | None = None) -> str:
    """Make text safe inside a Markdown table cell / link label."""
    s = " ".join(str(text or "").split())
    if limit and len(s) > limit:
        s = s[:limit - 1] + "…"
    return (s.replace("|", "\\|").replace("[", "(").replace("]", ")")
             .replace("<", "&lt;").replace(">", "&gt;"))


def _md_link(item, limit=50) -> str:
    title = md_cell(item.get("title") or "?", limit)
    return f"[{title}]({item['url']})" if item.get("url") else title


def select(items, max_price: float, max_bid: float):
    """Split listings into priced-within-budget per category, plus unpriced property."""
    categories = {cat: [] for cat in SECTION_NAMES}
    unknown_imoveis = []
    for item in items:
        price, bid = item.get("price"), item.get("current_bid") or 0
        if price is None:
            if item["category"] == "imoveis":
                unknown_imoveis.append(item)
            continue
        if price <= max_price and bid <= max_bid:
            categories[item["category"]].append(item)
    for cat in categories:
        categories[cat].sort(key=lambda it: (-it["score"], (it["current_bid"] or 0) / it["price"]
                                             if it["price"] else 999))
    unknown_imoveis.sort(key=lambda it: -it["score"])
    return categories, unknown_imoveis


def _table(lines, rows, *, with_country=False):
    head = "| # | Score | " + ("Country | " if with_country else "") + \
        "Title | Price | Bid | Location | Ends | Why |"
    lines.append(head)
    lines.append("|" + "---|" * (head.count("|") - 1))
    for i, it in enumerate(rows, 1):
        country = f"{it.get('country')} | " if with_country else ""
        drop = f" ↓{it['price_drop_pct']:.0f}%" if it.get("price_drop_pct") else ""
        lines.append(
            f"| {i} | {it['score']:.0f} | {country}{_md_link(it)} | {_money(it['price'])}{drop} | "
            f"{_money(it['current_bid']) if it.get('current_bid') else '-'} | {md_cell(_loc(it), 40)} | "
            f"{(it.get('date_end') or '-')[:10]} | {md_cell(', '.join(it['reasons']), 70)} |"
        )
    lines.append("")


def build_markdown(items, hidden_counts: Counter, health, max_price, max_bid, now) -> tuple[str, dict]:
    categories, unknown_imoveis = select(items, max_price, max_bid)
    # Score 0 = fractional share / usufruct: listed per country, never a "pick".
    props = [it for it in categories["imoveis"] if it["score"] > 0]

    hidden = ", ".join(f"{n} {k}" for k, n in hidden_counts.most_common()) or "none"
    lines = [
        "# EU Investment Scanner Report",
        f"**Generated**: {now:%Y-%m-%d %H:%M} UTC  ",
        f"**Budget**: €{max_price:,.0f}  ",
        f"**Listings shown**: {len(items)} (hidden: {hidden})  ",
        "",
    ]

    broken = [h for h in health if h["state"] in ("error", "broken")]
    if broken:
        lines.append(f"> ⚠️ {len(broken)} source(s) failing: "
                     + ", ".join(f"{h['source']} ({h['state']})" for h in broken)
                     + " — see *Source health* at the end.")
        lines.append("")

    lines.append("## Top picks — all countries")
    lines.append("")
    if props:
        _table(lines, props[:25], with_country=True)
    else:
        lines.append("_No priced property listings in budget._\n")

    soon = [it for it in props if (d := days_left(it.get("date_end"), now)) is not None and 0 < d <= 7]
    soon.sort(key=lambda it: it.get("date_end") or "")
    lines.append(f"## Ending within 7 days — {len(soon)}")
    lines.append("")
    if soon:
        _table(lines, soon[:30], with_country=True)
    else:
        lines.append("_Nothing in budget ends this week._\n")

    for cat, label in SECTION_NAMES.items():
        cat_items = categories[cat]
        lines.append(f"## {label} — {len(cat_items)} listings")
        lines.append("")
        if not cat_items and not (cat == "imoveis" and unknown_imoveis):
            lines.append("_None found._\n")
            continue

        by_country = defaultdict(list)
        for it in cat_items:
            by_country[it.get("country") or "PT"].append(it)
        for cc in _countries(cat_items):
            c_items = by_country[cc]
            lines.append(f"### {COUNTRY_NAMES.get(cc, cc)} ({len(c_items)})")
            lines.append("")
            _table(lines, c_items[:30] if cat == "imoveis" else c_items[:15])

        if cat == "imoveis" and unknown_imoveis:
            lines.append("### Price unknown — open to check")
            lines.append("")
            lines.append("_No published price; scored on everything else._")
            lines.append("")
            by_unknown = defaultdict(list)
            for it in unknown_imoveis:
                by_unknown[it.get("country") or "PT"].append(it)
            for cc in _countries(unknown_imoveis):
                u_items = by_unknown[cc]
                lines.append(f"#### {COUNTRY_NAMES.get(cc, cc)} ({len(u_items)})")
                lines.append("")
                lines.append("| # | Score | Title | Location | Ends |")
                lines.append("|---|---|---|---|---|")
                for i, it in enumerate(u_items[:15], 1):
                    lines.append(f"| {i} | {it['score']:.0f} | {_md_link(it)} | {md_cell(_loc(it), 40)} | "
                                 f"{(it.get('date_end') or '-')[:10]} |")
                lines.append("")

    lines.append("## Source health")
    lines.append("")
    lines.append("| Source | State | Last run | Last count | Failing runs | Last message |")
    lines.append("|---|---|---|---|---|---|")
    for h in health:
        lines.append(f"| {h['source']} | {h['state']} | {(h['last_run'] or '-')[:16]} | "
                     f"{h['last_count'] if h['last_count'] is not None else '-'} | {h['failing_runs']} | "
                     f"{md_cell(h['last_message'] or '', 80)} |")
    lines.append("")

    # LLM-ready compact summary for the top imóveis (saves tokens)
    lines.append("## Top Imóveis — Compact (for LLM analysis)")
    lines.append("```json")
    compact = [{
        "t": (it["title"] or "")[:60], "vb": it["price"], "bid": it["current_bid"],
        "loc": _loc(it), "tipo": it["tipo"], "area": it["area_m2"],
        "ends": (it["date_end"] or "")[:10], "url": it["url"], "score": it["score"],
    } for it in props[:15]]
    lines.append(json.dumps(compact, ensure_ascii=False, indent=1))
    lines.append("```")
    lines.append("")
    return "\n".join(lines), {"categories": categories, "unknown": unknown_imoveis}


def generate_report(db, max_price: float = 50000, max_bid: float | None = None, *,
                    filters: dict | None = None, out_dir: str | None = None,
                    desktop_copy: bool = True, known_sources=None) -> str:
    """Write report.md (+ report.docx and report.pdf when the libraries are
    installed) and return the Markdown path."""
    now = utcnow()
    max_bid = max_price if max_bid is None else max_bid
    out_dir = out_dir or HERE

    everything = load_listings(db, filters=filters, include_hidden=True, now=now)
    items = [it for it in everything if not it["hidden_reason"]]
    hidden_counts = Counter(hidden_category(it["hidden_reason"]) for it in everything
                            if it["hidden_reason"])
    health = source_health(db, known_sources)

    report, parts = build_markdown(items, hidden_counts, health, max_price, max_bid, now)
    report_path = os.path.join(out_dir, "report.md")
    with open(report_path, "w", encoding="utf-8") as f:
        f.write(report)
    cats = parts["categories"]
    LOG.info(f"Report: {len(cats['imoveis'])} imóveis, {len(cats['ouro_joias'])} ouro/joias, "
             f"{len(cats['outros'])} outros → {report_path}")

    desktop = os.path.join(os.path.expanduser("~"), "Desktop") if desktop_copy else None
    try:
        _write_docx(parts, now, max_price, len(items), out_dir, desktop)
    except ImportError:
        LOG.debug("python-docx not installed; skipping report.docx")
    except Exception as e:
        LOG.warning(f"Word report generation failed: {e}")
    try:
        _write_pdf(parts, now, max_price, len(items), out_dir, desktop)
    except ImportError:
        LOG.debug("fpdf2 not installed; skipping report.pdf")
    except Exception as e:
        LOG.warning(f"PDF generation failed: {e}")
    return report_path


def _save_copy(save, primary: str, desktop_dir: str | None, name: str):
    save(primary)
    if desktop_dir and os.path.isdir(desktop_dir):
        target = os.path.join(desktop_dir, name)
        try:
            save(target)
        except OSError:
            LOG.warning(f"Could not write {target} (file open?), skipping desktop copy")


def _add_hyperlink(paragraph, url: str, text: str, size_pt: int = 9):
    """python-docx has no public hyperlink API; this is the standard oxml recipe."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E99")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size_pt * 2))
    props.extend([color, underline, sz])
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _write_docx(parts, now: datetime, max_price, n_items, out_dir, desktop):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    doc.add_heading("EU Investment Scanner Report", level=0)
    doc.add_paragraph(f"Generated: {now:%Y-%m-%d %H:%M} UTC    Budget: €{max_price:,.0f}    Listings: {n_items}")

    for cat, label in SECTION_NAMES.items():
        cat_items = parts["categories"][cat]
        doc.add_heading(f"{label} — {len(cat_items)} listings", level=1)
        if not cat_items:
            doc.add_paragraph("None found.")
            continue
        by_country = defaultdict(list)
        for it in cat_items:
            by_country[it.get("country") or "PT"].append(it)
        for cc in _countries(cat_items):
            c_items = by_country[cc]
            doc.add_heading(f"{COUNTRY_NAMES.get(cc, cc)} ({len(c_items)})", level=2)
            table = doc.add_table(rows=1, cols=6)
            table.style = "Light Grid Accent 1"
            for i, hdr in enumerate(["#", "Score", "Title", "Price", "Location", "Why"]):
                table.rows[0].cells[i].text = hdr
            for idx, it in enumerate(c_items[:30] if cat == "imoveis" else c_items[:15], 1):
                row = table.add_row().cells
                row[0].text = str(idx)
                row[1].text = f"{it['score']:.0f}"
                title = (it["title"] or "?")[:60]
                if it.get("url"):
                    _add_hyperlink(row[2].paragraphs[0], it["url"], title)
                else:
                    row[2].text = title
                row[3].text = _money(it["price"])
                row[4].text = _loc(it)[:30]
                row[5].text = ", ".join(it["reasons"])[:60]

    _save_copy(doc.save, os.path.join(out_dir, "report.docx"), desktop, "Auction-Report.docx")


def _pdf_font(pdf):
    """DejaVu (full Unicode) if cartas.py already downloaded it, else Helvetica."""
    font_dir = os.path.join(HERE, "fonts")
    regular = os.path.join(font_dir, "DejaVuSans.ttf")
    bold = os.path.join(font_dir, "DejaVuSans-Bold.ttf")
    if os.path.exists(regular) and os.path.exists(bold):
        pdf.add_font("DejaVu", "", regular)
        pdf.add_font("DejaVu", "B", bold)
        return "DejaVu", (lambda t: t or "")
    return "Helvetica", (lambda t: (t or "").encode("latin-1", "replace").decode("latin-1"))


def _write_pdf(parts, now: datetime, max_price, n_items, out_dir, desktop):
    from fpdf import FPDF

    pdf = FPDF(orientation="L", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    font, s = _pdf_font(pdf)
    pdf.set_font(font, "B", 16)
    pdf.cell(0, 10, "EU Investment Scanner Report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(font, "", 9)
    pdf.cell(0, 6, s(f"Generated: {now:%Y-%m-%d %H:%M} UTC    Budget: EUR {max_price:,.0f}    "
                     f"Listings: {n_items}    (titles are links)"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    col_w = [8, 12, 105, 25, 50, 77]
    for cat, label in SECTION_NAMES.items():
        cat_items = parts["categories"][cat]
        pdf.set_font(font, "B", 12)
        pdf.cell(0, 8, s(f"{label} - {len(cat_items)} listings"), new_x="LMARGIN", new_y="NEXT")
        if not cat_items:
            pdf.set_font(font, "", 9)
            pdf.cell(0, 6, "None found.", new_x="LMARGIN", new_y="NEXT")
            continue
        by_country = defaultdict(list)
        for it in cat_items:
            by_country[it.get("country") or "PT"].append(it)
        for cc in _countries(cat_items):
            c_items = by_country[cc]
            pdf.set_font(font, "B", 10)
            pdf.cell(0, 7, s(f"{COUNTRY_NAMES.get(cc, cc)} ({len(c_items)})"), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font(font, "B", 7)
            for w, h in zip(col_w, ["#", "Score", "Title", "Price", "Location", "Why"]):
                pdf.cell(w, 5, h, border=1)
            pdf.ln()
            pdf.set_font(font, "", 7)
            for idx, it in enumerate(c_items[:30] if cat == "imoveis" else c_items[:15], 1):
                cells = [
                    str(idx), f"{it['score']:.0f}", s((it["title"] or "?")[:70]),
                    f"EUR {it['price']:,.0f}" if it["price"] else "?",
                    s(_loc(it)[:32]), s(", ".join(it["reasons"])[:55]),
                ]
                for j, (w, c) in enumerate(zip(col_w, cells)):
                    pdf.cell(w, 4, c, border=1, link=(it.get("url") or "") if j == 2 else "")
                pdf.ln()
            pdf.ln(3)

    _save_copy(pdf.output, os.path.join(out_dir, "report.pdf"), desktop, "Auction-Report.pdf")


def _safe_print(text: str):
    try:
        print(text)
    except UnicodeEncodeError:
        print(text.encode("ascii", errors="replace").decode("ascii"))


def print_console_summary(db, max_price: float = 50000, *, filters: dict | None = None,
                          items=None, known_sources=None):
    """Top 5 property picks per country, then any failing sources."""
    now = utcnow()
    items = items if items is not None else load_listings(db, filters=filters, now=now)
    props = [it for it in items if it["category"] == "imoveis"
             and it.get("price") is not None and it["price"] <= max_price]
    total_by_country = Counter(it.get("country") or "PT" for it in items)

    _safe_print(f"\n{'=' * 60}")
    _safe_print(f"  EU AUCTION SCANNER - {now:%d %b %Y %H:%M} UTC")
    _safe_print(f"  {len(items)} listings across {len(total_by_country)} countries (budget EUR {max_price:,.0f})")
    _safe_print(f"{'=' * 60}")

    by_country = defaultdict(list)
    for it in props:
        by_country[it.get("country") or "PT"].append(it)
    for cc in _countries(props):
        top5 = sorted(by_country[cc], key=lambda it: -it["score"])[:5]
        _safe_print(f"\n  {COUNTRY_NAMES.get(cc, cc)} ({total_by_country[cc]} total)")
        _safe_print(f"  {'-' * 56}")
        for i, it in enumerate(top5, 1):
            _safe_print(f"  {i}. [{it['score']:.0f}] {(it['title'] or '?')[:40]}")
            _safe_print(f"     EUR {it['price']:,.0f}  {_loc(it)[:20]}  {(it.get('date_end') or '')[:10]}")
            _safe_print(f"     {', '.join(it['reasons'])[:60]}")
            _safe_print(f"     {it.get('url') or ''}")

    failing = [h for h in source_health(db, known_sources) if h["state"] in ("error", "broken")]
    if failing:
        _safe_print(f"\n  Sources needing attention ({len(failing)}):")
        for h in failing:
            msg = f" — {h['last_message'][:70]}" if h["last_message"] else ""
            _safe_print(f"   ! {h['source']}: {h['state']}, {h['failing_runs']} run(s) without results{msg}")
    _safe_print(f"\n{'=' * 60}\n")


def print_sealed_bid_summary(db, *, filters: dict | None = None):
    from scoring import SEALED_BID_PATTERNS
    from common import has_term

    items = load_listings(db, filters=filters, where="source = 'citius'")
    sealed = [it for it in items
              if has_term(f"{it.get('title') or ''} {it.get('description') or ''}",
                          SEALED_BID_PATTERNS, negations=False)]
    sealed.sort(key=lambda it: -it["score"])
    print(f"\n{'=' * 60}")
    print(f"  SEALED-BID LISTINGS (carta fechada) — {len(sealed)} found")
    print(f"{'=' * 60}")
    for it in sealed[:20]:
        raw = json.loads(it.get("raw_json") or "{}")
        _safe_print(f"\n  [{it['score']:.0f}] {(it.get('title') or '?')[:60]}")
        _safe_print(f"       Price: {_money(it.get('price'))}")
        _safe_print(f"       Location: {_loc(it)}")
        _safe_print(f"       Ends: {(it.get('date_end') or '')[:10]}")
        _safe_print(f"       Agente: {raw.get('agente_nome', '')}  Tel: {raw.get('agente_contacto', '')}")
        _safe_print(f"       Email: {raw.get('agente_email', '')}")
        _safe_print(f"       Flags: {', '.join(it['reasons'])}")
    print()
