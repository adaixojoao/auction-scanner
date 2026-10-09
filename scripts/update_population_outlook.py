"""Refresh the EUROPOP2019 regional population scenario used as a small caution.

    python scripts/update_population_outlook.py

Needs openpyxl (pip install openpyxl) and a network connection. Writes
data/europop2019.csv and data/place_nuts3.csv. The app only reads those files.

The projection is Eurostat's EUROPOP2019 baseline (proj_19rp3): population on
1 January 2019 and 2050 for each NUTS 3 region. Eurostat presents it as one
scenario, not a forecast. A municipality is attached to that region by
Eurostat's LAU 2021 correspondence (the name of the commune, or, in Portugal,
the concelho: the first four digits of the parish code). A name that sits in
two regions is left out.
"""
from __future__ import annotations

import csv
import io
import json
import os
import sys
import urllib.request

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)

from prices import place_key  # noqa: E402

PROJECTION_URL = (
    "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/proj_19rp3"
    "?format=JSON&lang=en&projection=BSL&sex=T&age=TOTAL&time=2019&time=2050")
LAU_URL = "https://ec.europa.eu/eurostat/documents/345175/501971/EU-27-LAU-2021-NUTS-2021.xlsx"
# Official concelho codes (district + concelho). The same four digits start
# every parish code in the Eurostat LAU table, which carries the NUTS 3 code.
PT_MUNICIPALITIES_URL = (
    "https://raw.githubusercontent.com/centraldedados/codigos_postais/master/data/concelhos.csv")
OUT_REGIONS = os.path.join(HERE, "data", "europop2019.csv")
OUT_PLACES = os.path.join(HERE, "data", "place_nuts3.csv")
# Eurostat's Greece code. The scanner uses GR.
COUNTRY = {"EL": "GR"}


def _get(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "auction-scanner"})
    with urllib.request.urlopen(req, timeout=120) as response:
        return response.read()


def _regions(payload: dict) -> dict[str, dict]:
    dims = payload["dimension"]
    geos = list(dims["geo"]["category"]["index"])
    times = list(dims["time"]["category"]["index"])
    labels = dims["geo"]["category"]["label"]
    values = payload["value"]
    # size is [freq, projection, age, sex, unit, geo, time]; only geo and time vary.
    n_time = len(times)
    out = {}
    for geo_i, code in enumerate(geos):
        pair = []
        for time_i in range(n_time):
            raw = values.get(str(geo_i * n_time + time_i))
            pair.append(float(raw) if raw is not None else None)
        pop_2019, pop_2050 = (pair + [None, None])[:2]
        if not pop_2019 or not pop_2050:
            continue
        name = labels.get(code, code)
        name = name.replace(" (NUTS 2021)", "").strip()
        out[code] = {"name": name, "pop_2019": int(pop_2019), "pop_2050": int(pop_2050),
                     "change_pct": round((pop_2050 - pop_2019) / pop_2019 * 100, 1)}
    return out


def _places(regions: dict[str, dict]) -> dict[tuple[str, str], str]:
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(_get(LAU_URL)), read_only=True, data_only=True)
    grouped: dict[tuple[str, str], set[str]] = {}
    pt_nuts: dict[str, set[str]] = {}
    for sheet in book.sheetnames:
        if len(sheet) != 2 or sheet in ("Fi",):
            continue
        country = COUNTRY.get(sheet, sheet)
        rows = book[sheet].iter_rows(values_only=True)
        header = [str(cell or "") for cell in next(rows)]
        try:
            nuts_at = header.index("NUTS 3 CODE")
            latin_at = header.index("LAU NAME LATIN")
            national_at = header.index("LAU NAME NATIONAL")
            code_at = header.index("LAU CODE")
        except ValueError:
            continue
        for row in rows:
            nuts = str(row[nuts_at] or "").strip()
            if nuts not in regions:
                continue
            if country == "PT":
                code = str(row[code_at] or "").strip()
                if len(code) >= 4:
                    pt_nuts.setdefault(code[:4], set()).add(nuts)
            for col in (latin_at, national_at):
                key = place_key(str(row[col] or ""))
                if len(key) < 3:
                    continue
                grouped.setdefault((country, key), set()).add(nuts)
    places = {key: next(iter(nuts)) for key, nuts in grouped.items() if len(nuts) == 1}
    # Portugal's listings name the concelho, which is often not itself a parish.
    text = _get(PT_MUNICIPALITIES_URL).decode("utf-8")
    for row in csv.DictReader(io.StringIO(text)):
        code = f"{int(row['cod_distrito']):02d}{int(row['cod_concelho']):02d}"
        nuts = pt_nuts.get(code) or set()
        key = place_key(row.get("nome_concelho") or "")
        if len(nuts) == 1 and len(key) >= 3:
            places[("PT", key)] = next(iter(nuts))
    return places


def main() -> None:
    regions = _regions(json.loads(_get(PROJECTION_URL)))
    places = _places(regions)
    with open(OUT_REGIONS, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["nuts3", "name", "pop_2019", "pop_2050", "change_pct"])
        for code in sorted(regions):
            row = regions[code]
            writer.writerow([code, row["name"], row["pop_2019"], row["pop_2050"], row["change_pct"]])
    with open(OUT_PLACES, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["country", "name_key", "nuts3"])
        for (country, key), nuts in sorted(places.items()):
            writer.writerow([country, key, nuts])
    steep = sorted((row["change_pct"], code, row["name"]) for code, row in regions.items()
                   if code.startswith("PT"))
    print(f"{len(regions)} regions, {len(places)} place names")
    print("Portugal, smallest to largest change:")
    for change, code, name in steep:
        print(f"  {change:6.1f}%  {code}  {name}")


if __name__ == "__main__":
    main()
