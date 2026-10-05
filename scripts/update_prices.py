"""
Refresh data/pt_home_prices.csv: the median price per m² of homes sold in every
Portuguese municipality, from Statistics Portugal (INE), for the score's
"X% below local prices".

    python scripts/update_prices.py                      # the default indicator
    python scripts/update_prices.py --indicator 0012345  # another INE indicator code

Run it on a PC that can reach www.ine.pt, check the summary it prints, and
commit data/pt_home_prices.csv through a pull request. INE publishes the
figures every quarter; refreshing once or twice a year is plenty.

The indicator wanted is INE's median sale value per m² of family dwellings by
municipality, quarterly, from the housing price statistics at local level. The
default, 0012234, is "Valor mediano das vendas de alojamentos familiares nos
últimos 12 meses (Metodologia 2022 - €/m²) por Localização geográfica (NUTS -
2024) e Categoria": published every quarter, each value the median of the last
12 months. It is the series that covers all 308 municipalities; INE's plain
quarterly median exists only for regions and cities over 100,000 people. If
INE renumbers it, find the code on ine.pt (Estatísticas → Indicadores, search
"valor mediano das vendas") and pass it with --indicator.
"""
from __future__ import annotations

import argparse
import csv
import io
import os
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)

from prices import COLUMNS, PARISH_COLUMNS, PRICE_FILES, PT_FILE, PT_PARISH_FILE, PT_RENT_FILE, RENT_FILES  # noqa: E402

API = "https://www.ine.pt/ine/json_indicador/pindica.jsp"
DEFAULT_INDICATOR = "0012234"
# `--rents`: "Valor mediano das rendas de novos contratos de arrendamento de
# alojamentos familiares nos últimos 12 meses (€/m²) por Localização geográfica",
# every six months, into data/pt_rents.csv. INE keeps small municipalities with
# too few leases secret, so there are fewer rows than for the sale prices.
RENT_INDICATOR = "0012598"
# `--rents-fr`: France's "carte des loyers" (Ministère du Logement / ANIL), the
# predicted rent per m² of a house in every commune, from data.gouv.fr.
# `--rents-es FILE`: Spain's Sistema Estatal de Referencia del Precio del Alquiler
# de Vivienda (SERPAVI), the "BD Sistema Estatal Índices de Alquiler de Vivienda"
# xlsx from mivau.gob.es/vivienda/alquila-bien-es-tu-derecho/serpavi (download it
# in a browser). Median rent per m² a month from income-tax returns.
FR_RENTS_DATASET = "https://www.data.gouv.fr/api/1/datasets/?q=carte%20des%20loyers%20par%20commune&page_size=20"


def parse_ine(payload, digits: int = 0) -> tuple[list[dict], str, str]:
    """(rows, period, indicator title) from INE's JSON API answer: the latest
    period's values for municipalities, for all kinds of dwelling ("Total")
    when the indicator splits them.

    Municipalities have 7-character codes. Since the 2024 regions (NUTS 2024)
    many contain letters ("11D1818" Sernancelhe, "1C20204" Barrancos); only
    reading all-digit codes kept 144 of the 308."""
    entry = payload[0] if isinstance(payload, list) else payload
    title = entry.get("IndicadorDsg", "")
    data = entry.get("Dados") or {}
    if not data:
        raise ValueError(f"no data in INE's answer: {str(entry)[:300]}")
    period = entry.get("UltimoPref") if entry.get("UltimoPref") in data else list(data)[-1]
    rows = []
    for rec in data[period]:
        code = str(rec.get("geocod", ""))
        if not (len(code) == 7 and code.isalnum()):          # municipalities only
            continue
        extra = [v for k, v in rec.items() if k.startswith("dim_") and k.endswith("_t")]
        if extra and not all(str(v).strip().lower() in ("total", "t") for v in extra):
            continue
        try:
            value = float(str(rec.get("valor", "")).replace(",", "."))
        except ValueError:
            continue
        rows.append({"municipality": rec.get("geodsg", "").strip(), "eur_m2": round(value, digits or None),
                     "period": period, "source": "INE"})
    return rows, period, title


def parse_ine_parishes(payload) -> list[dict]:
    """The parishes INE gives a figure for (9-character codes: the Porto and
    Lisbon areas, Setúbal, the Algarve and cities over 100,000 people), each
    under its municipality (the first 7 characters of its code)."""
    entry = payload[0] if isinstance(payload, list) else payload
    data = entry.get("Dados") or {}
    if not data:
        return []
    period = entry.get("UltimoPref") if entry.get("UltimoPref") in data else list(data)[-1]
    towns = {str(r.get("geocod")): r.get("geodsg", "").strip() for r in data[period]
             if len(str(r.get("geocod", ""))) == 7}
    rows = []
    for rec in data[period]:
        code = str(rec.get("geocod", ""))
        extra = [v for k, v in rec.items() if k.startswith("dim_") and k.endswith("_t")]
        if len(code) != 9 or code[:7] not in towns or (
                extra and not all(str(v).strip().lower() in ("total", "t") for v in extra)):
            continue
        try:
            value = float(str(rec.get("valor", "")).replace(",", "."))
        except ValueError:
            continue
        rows.append({"municipality": towns[code[:7]], "parish": rec.get("geodsg", "").strip(),
                     "eur_m2": round(value), "period": period, "source": "INE"})
    return sorted(rows, key=lambda r: (r["municipality"], r["parish"]))


def update_rents(requests) -> int:
    r = requests.get(API, params={"op": "2", "varcd": RENT_INDICATOR, "lang": "PT"}, timeout=120)
    r.raise_for_status()
    rows, period, title = parse_ine(r.json(), digits=2)
    print(f"Indicator {RENT_INDICATOR}: {title}")
    print(f"Period: {period} — {len(rows)} municipalities")
    if len(rows) < 150:
        print("Fewer than 150 municipalities: probably not the right indicator. Nothing written.")
        return 1
    rows.sort(key=lambda row: row["municipality"])
    with open(PT_RENT_FILE, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Written to {PT_RENT_FILE}")
    return 0


def parse_fr_rents(text: str, year: str) -> list[dict]:
    """The carte des loyers CSV (";"-separated, decimal commas, LIBGEO and
    loypredm2 columns) → rows; a commune name found twice keeps its first figure."""
    rows, seen = [], set()
    for rec in csv.DictReader(text.splitlines(), delimiter=";"):
        name = (rec.get("LIBGEO") or "").strip()
        try:
            value = float((rec.get("loypredm2") or "").replace(",", "."))
        except ValueError:
            continue
        if not name or value <= 0 or name in seen:
            continue
        seen.add(name)
        rows.append({"municipality": name, "eur_m2": round(value, 2), "period": year,
                     "source": "Carte des loyers"})
    return sorted(rows, key=lambda r: r["municipality"])


def update_rents_fr(requests) -> int:
    import re
    found = []
    for ds in requests.get(FR_RENTS_DATASET, timeout=60).json().get("data", []):
        year = re.search(r"par commune en (\d{4})", ds.get("title", ""))
        house = next((r["url"] for r in ds.get("resources", []) if "maison" in r.get("title", "").lower()
                      and r["url"].endswith(".csv")), None)
        if year and house:
            found.append((year.group(1), house))
    if not found:
        print("No carte des loyers found on data.gouv.fr. Nothing written.")
        return 1
    year, url = max(found)
    r = requests.get(url, timeout=120)
    r.raise_for_status()
    rows = parse_fr_rents(r.content.decode("latin-1"), year)
    print(f"Carte des loyers {year}: {len(rows)} communes")
    if len(rows) < 20000:
        print("Fewer than 20,000 communes: probably not the right file. Nothing written.")
        return 1
    with open(RENT_FILES["FR"], "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Written to {RENT_FILES['FR']}")
    return 0


FR_PRICES_DATASET = "https://www.data.gouv.fr/api/1/datasets/?q=Indicateurs%20Immobiliers%20par%20commune%20et%20par%20ann%C3%A9e&page_size=5"
FR_COMMUNES = "https://geo.api.gouv.fr/communes?fields=nom,code,codeDepartement&format=json"
FR_MIN_SALES = 5            # fewer sales in the year: the département's average instead


def fr_price_rows(text: str, communes: dict[str, tuple[str, str]], year: str) -> list[dict]:
    """DVF's average €/m² per commune (INSEE code) → rows by name, by "name|département",
    and a sales-weighted average per département ("prov:<code>")."""
    by_name: dict[str, tuple[int, dict]] = {}
    rows, dept = [], {}
    for rec in csv.DictReader(io.StringIO(text)):
        try:
            sales, eur_m2 = int(rec["nb_mutations"]), float(rec["Prixm2Moyen"])
        except (KeyError, TypeError, ValueError):
            continue
        found = communes.get(rec.get("INSEE_COM", ""))
        if not found or eur_m2 <= 0:
            continue
        name, code = found
        total = dept.setdefault(code, [0.0, 0])
        total[0] += eur_m2 * sales
        total[1] += sales
        if sales < FR_MIN_SALES:
            continue
        row = {"municipality": name, "eur_m2": round(eur_m2), "period": year, "source": "DVF"}
        rows.append({**row, "municipality": f"{name}|{code}"})
        if sales > by_name.get(name, (0, None))[0]:          # same-named communes: the busier one
            by_name[name] = (sales, row)
    rows += [row for _, row in by_name.values()]
    rows += [{"municipality": f"prov:{code}", "eur_m2": round(s / n), "period": year, "source": "DVF"}
             for code, (s, n) in dept.items() if n]
    return rows


def update_prices_fr(requests) -> int:
    import re
    found = []
    for ds in requests.get(FR_PRICES_DATASET, timeout=60).json().get("data", []):
        for res in ds.get("resources", []):
            year = re.search(r"(20\d\d)", res.get("title", ""))
            if year and res.get("format") == "csv" and "DVF_Communes" in res.get("title", ""):
                found.append((year.group(1), res["url"]))
    if not found:
        print("No DVF commune indicators found on data.gouv.fr. Nothing written.")
        return 1
    year, url = max(found)
    communes = {c["code"]: (c["nom"], c["codeDepartement"])
                for c in requests.get(FR_COMMUNES, timeout=120).json()}
    r = requests.get(url, timeout=180)
    r.raise_for_status()
    rows = fr_price_rows(r.content.decode("utf-8", "replace"), communes, year)
    print(f"DVF {year}: {len(rows)} rows")
    if len(rows) < 10000:
        print("Fewer than 10,000 rows: probably not the right file. Nothing written.")
        return 1
    return _write(PRICE_FILES["FR"], rows)


NL_CBS = "https://opendata.cbs.nl/ODataApi/odata/83625NED"
NL_HOME_M2 = 120            # CBS gives the average price per home, not per m²: an average home's size


def update_prices_nl(requests) -> int:
    """CBS: the average sale price of existing homes per gemeente and province,
    the latest full year, ÷ NL_HOME_M2."""
    regions = {r["Key"].strip(): r["Title"] for r in requests.get(f"{NL_CBS}/RegioS", timeout=60).json()["value"]}
    year = max(p["Key"] for p in requests.get(f"{NL_CBS}/Perioden", timeout=60).json()["value"]
               if p["Key"].endswith("JJ00"))
    data = requests.get(f"{NL_CBS}/TypedDataSet", params={"$filter": f"Perioden eq '{year}'"}, timeout=120).json()
    rows = []
    for rec in data.get("value", []):
        key, price = rec["RegioS"].strip(), rec.get("GemiddeldeVerkoopprijs_1")
        if not price or key not in regions:
            continue
        name = regions[key]
        eur_m2 = round(price / NL_HOME_M2)
        source = f"CBS average price ÷ {NL_HOME_M2} m²"
        if key.startswith("GM"):
            rows.append({"municipality": name, "eur_m2": eur_m2, "period": year[:4], "source": source})
        elif key.startswith("PV"):
            rows.append({"municipality": "prov:" + name.replace(" (PV)", ""), "eur_m2": eur_m2,
                         "period": year[:4], "source": source})
    print(f"CBS {year[:4]}: {len(rows)} gemeenten and provinces")
    if len(rows) < 300:
        print("Fewer than 300 rows: probably not the right table. Nothing written.")
        return 1
    return _write(PRICE_FILES["NL"], rows)


LU_DATASET = "https://data.public.lu/api/1/datasets/?q=prix%20annonc%C3%A9s%20commune&page_size=3"


def update_prices_lu(requests) -> int:
    """Observatoire de l'Habitat: asking prices of houses per m² per commune, the
    latest year of the series (communes with under 30 ads have no figure)."""
    import openpyxl
    ds = requests.get(LU_DATASET, timeout=60).json()["data"][0]
    url = next(r["url"] for r in ds["resources"] if r["format"] == "xlsx" and "maisons" in r["title"].lower())
    r = requests.get(url, timeout=120)
    r.raise_for_status()
    sheet = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True).worksheets[-1]
    rows = []
    for rec in sheet.iter_rows(values_only=True):
        name, eur_m2 = (rec[2], rec[5]) if len(rec) > 5 else (None, None)
        if isinstance(name, str) and isinstance(eur_m2, (int, float)) and eur_m2 > 0:
            rows.append({"municipality": name.strip(), "eur_m2": round(eur_m2), "period": sheet.title,
                         "source": "Observatoire de l'Habitat (asking)"})
    print(f"Luxembourg {sheet.title}: {len(rows)} communes")
    if len(rows) < 40:
        print("Fewer than 40 communes: probably not the right file. Nothing written.")
        return 1
    return _write(PRICE_FILES["LU"], rows)


ES_TOWNS = "https://apps.fomento.gob.es/BoletinOnline2/sedal/35103500.XLS"      # towns over 25,000 people
ES_PROVINCES = "https://apps.fomento.gob.es/BoletinOnline2/sedal/35101000.XLS"  # provinces, quarterly
_ES_ARTICLES = {"a", "o", "el", "la", "los", "las", "les", "els", "illes", "l'"}


def es_name(name: str) -> str:
    """The ministry's "Coruña (A)" / "Ejido (El)" / "Balears (Illes)" → "A Coruña" / "El Ejido" /
    "Illes Balears"; "Asturias (Principado de)" → "Asturias"."""
    import re
    name = " ".join(str(name).split())
    m = re.fullmatch(r"(.+?)\s*\((.+?)\)", name)
    if not m:
        return name
    return f"{m.group(2)} {m.group(1)}" if m.group(2).lower() in _ES_ARTICLES else m.group(1)


def _number(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def update_prices_es(requests) -> int:
    """MIVAU: the appraised value of homes per m² — per town (over 25,000 people),
    and per province and region for everywhere else ("prov:<name>")."""
    import xlrd
    towns = xlrd.open_workbook(file_contents=requests.get(ES_TOWNS, timeout=180).content).sheets()[-1]
    period = next((str(v).replace("(*)", "").strip() for r in range(towns.nrows) for v in towns.row_values(r)[1:2]
                   if "trimestre" in str(v).lower()), towns.name)
    rows = []
    for r in range(towns.nrows):
        values = towns.row_values(r)
        town, total = (values[2], _number(values[5])) if len(values) > 5 else (None, None)
        if isinstance(town, str) and town.strip() and total and total > 0:
            rows.append({"municipality": es_name(town), "eur_m2": round(total), "period": period,
                         "source": "MIVAU appraised"})
    provinces = xlrd.open_workbook(file_contents=requests.get(ES_PROVINCES, timeout=180).content).sheets()[-1]
    for r in range(provinces.nrows):
        values = provinces.row_values(r)
        name = values[1] if len(values) > 3 else None
        figures = [v for v in (_number(x) for x in values[2:-2]) if v]
        if isinstance(name, str) and name.strip() and figures and "total" not in name.lower():
            rows.append({"municipality": "prov:" + es_name(name), "eur_m2": round(figures[-1]),
                         "period": period, "source": "MIVAU appraised"})
    print(f"MIVAU {period}: {len(rows)} towns, provinces and regions")
    if len(rows) < 200:
        print("Fewer than 200 rows: probably not the right files. Nothing written.")
        return 1
    return _write(PRICE_FILES["ES"], rows)


def _write(path: str, rows: list[dict]) -> int:
    rows.sort(key=lambda row: row["municipality"])
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Written to {path}")
    return 0


def serpavi_rows(header: list, rows, name_col: str, prefix: str = "") -> list[dict]:
    """One row per place from a SERPAVI sheet: the latest year's median rent per
    m² of flats (ALQM2_LV_M_VC_yy), else of houses (…_VU_yy)."""
    cols = {h: i for i, h in enumerate(header) if h}
    years = sorted({h[-2:] for h in cols if str(h).startswith("ALQM2_LV_M_V")}, reverse=True)
    out = []
    for row in rows:
        name = str(row[cols[name_col]] or "").strip()
        if not name:
            continue
        for yy in years:
            value = next((row[cols[c]] for c in (f"ALQM2_LV_M_VC_{yy}", f"ALQM2_LV_M_VU_{yy}")
                          if c in cols and isinstance(row[cols[c]], (int, float)) and row[cols[c]] > 0), None)
            if value:
                out.append({"municipality": prefix + name, "eur_m2": round(float(value), 2),
                            "period": f"20{yy}", "source": "SERPAVI"})
                break
    return out


def update_rents_es(path: str) -> int:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = []
    for sheet, name_col, prefix in (("Municipios", "NMUN", ""), ("Provincias", "LITPRO", "prov:")):
        it = wb[sheet].iter_rows(values_only=True)
        header = list(next(it))
        rows += serpavi_rows(header, it, name_col, prefix)
    towns = sum(1 for r in rows if not r["municipality"].startswith("prov:"))
    print(f"SERPAVI: {towns} municipios, {len(rows) - towns} provinces")
    if towns < 2000:
        print("Fewer than 2,000 municipios: probably not the right file. Nothing written.")
        return 1
    with open(RENT_FILES["ES"], "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Written to {RENT_FILES['ES']}")
    return 0


def main(argv=None) -> int:
    import requests

    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--indicator", default=DEFAULT_INDICATOR)
    ap.add_argument("--out", default=PT_FILE)
    ap.add_argument("--rents", action="store_true", help="the monthly rents per m², into data/pt_rents.csv")
    ap.add_argument("--rents-fr", action="store_true", help="France's rents per m² per commune, into data/fr_rents.csv")
    ap.add_argument("--rents-es", metavar="XLSX", help="Spain's SERPAVI workbook → data/es_rents.csv")
    ap.add_argument("--prices-fr", action="store_true", help="France's €/m² per commune (DVF) → data/fr_home_prices.csv")
    ap.add_argument("--prices-nl", action="store_true", help="the Netherlands' prices per gemeente (CBS) → data/nl_home_prices.csv")
    ap.add_argument("--prices-lu", action="store_true", help="Luxembourg's asking €/m² per commune → data/lu_home_prices.csv")
    ap.add_argument("--prices-es", action="store_true", help="Spain's appraised €/m² (MIVAU; needs xlrd) → data/es_home_prices.csv")
    args = ap.parse_args(argv)
    if args.rents:
        return update_rents(requests)
    if args.rents_fr:
        return update_rents_fr(requests)
    if args.rents_es:
        return update_rents_es(args.rents_es)
    if args.prices_fr:
        return update_prices_fr(requests)
    if args.prices_nl:
        return update_prices_nl(requests)
    if args.prices_lu:
        return update_prices_lu(requests)
    if args.prices_es:
        return update_prices_es(requests)

    r = requests.get(API, params={"op": "2", "varcd": args.indicator, "lang": "PT"}, timeout=60)
    r.raise_for_status()
    rows, period, title = parse_ine(r.json())
    print(f"Indicator {args.indicator}: {title}")
    print(f"Period: {period} — {len(rows)} municipalities")
    if len(rows) < 250:
        print("Fewer than 250 municipalities: probably not the right indicator. Nothing written.")
        return 1
    rows.sort(key=lambda row: row["municipality"])
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with open(args.out, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    cheapest = sorted(rows, key=lambda row: row["eur_m2"])[:3]
    dearest = sorted(rows, key=lambda row: -row["eur_m2"])[:3]
    print("Cheapest:", ", ".join(f"{x['municipality']} €{x['eur_m2']}/m²" for x in cheapest))
    print("Dearest: ", ", ".join(f"{x['municipality']} €{x['eur_m2']}/m²" for x in dearest))
    print(f"Written to {args.out}")
    parishes = parse_ine_parishes(r.json())
    if parishes and args.out == PT_FILE:
        with open(PT_PARISH_FILE, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=PARISH_COLUMNS)
            writer.writeheader()
            writer.writerows(parishes)
        print(f"{len(parishes)} parishes written to {PT_PARISH_FILE}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
