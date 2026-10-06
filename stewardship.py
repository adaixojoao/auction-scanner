"""stewardship.py — a non-binding Climate & Land Stewardship Plan for rural land.

Built only from what the scanner already has (listing size and kind, geo
location confidence, climate.py layers via scoring.climate_score). It never
claims site-specific ecology, water rights, protected-area status or planting
permission. Opportunities that name cork, oak or riparian work are labelled
"verify locally". Cost figures are planning allowances, not quotes.
"""
from __future__ import annotations

from common import utcnow_iso
from scoring import climate_score, floods, property_kind

RURAL_KINDS = ("rural_plot",)
DATA_NOTE = ("Public climate maps and the sale notice — municipal-to-grid scale. "
             "Not a survey, not an ecological assessment, not legal advice.")


def policy(settings: dict | None) -> dict:
    from config import DEFAULTS
    return {**DEFAULTS.get("stewardship", {}), **((settings or {}).get("stewardship") or {})}


def available(item: dict, settings: dict | None = None) -> tuple[bool, str | None]:
    """Whether this listing gets a plan. Rural by default; mixed only if Settings says so."""
    kind = item.get("kind") or property_kind(item)
    if kind in RURAL_KINDS:
        return True, None
    if policy(settings).get("enable_for_mixed") and kind in ("home", "urban_plot", None):
        # Explicit opt-in for a house-with-land or unknown: still show with a warning.
        return True, "Enabled in Settings for non-rural listings; treat opportunities as even more tentative."
    return False, "Only shown for rural land (or turn on Settings → Stewardship for mixed properties)."


def _ha(item: dict) -> float | None:
    area = item.get("area_m2") or 0
    return area / 10000 if area else None


def _climate(item: dict) -> dict:
    return climate_score(item.get("climate"), item.get("kind") or property_kind(item))


def _location(item: dict) -> dict:
    import geo
    return geo.location_confidence(item)


def build(item: dict, settings: dict | None = None) -> dict:
    """Structured plan, or {"available": False, "reason": ...}."""
    ok, why = available(item, settings)
    if not ok:
        return {"available": False, "reason": why}

    p = policy(settings)
    clim = _climate(item)
    loc = _location(item)
    raw = item.get("climate") or {}
    flags = set(clim.get("flags") or [])
    ha = _ha(item)
    country = (item.get("country") or "PT").upper()
    profile = p.get("profile") or {}
    region = profile.get("region_label") or "your region"
    unknowns: list[str] = []
    sources = [
        {"label": "Climate layers (heat, hot days, water stress, fire, flood)", "when": "as of last scan"},
        {"label": "Location confidence", "when": loc.get("label") or "unknown"},
        {"label": "Plan generated", "when": utcnow_iso()[:10]},
    ]

    if clim.get("grade") == "unknown":
        unknowns.append("climate layers not read for this place yet")
    if loc.get("level") not in ("exact", "street", "cadastre", "verified"):
        unknowns.append(f"exact position ({(loc.get('label') or 'unknown').lower()})")
    if not ha:
        unknowns.append("land area")

    loc_warn = None
    if loc.get("level") in ("municipality", "unknown") or clim.get("confidence") != "exact":
        loc_warn = (f"{loc.get('label') or 'Location unknown'}: climate and water/fire/flood figures "
                    "may describe the town or parish, not this parcel. Verify the position "
                    "(Listings → ⓘ → Verify location) before relying on any of this.")

    suitability = _suitability(clim, flags, ha, loc_warn)
    risks = _risks(clim, flags, raw)
    opportunities = _opportunities(flags, raw, country, profile, floods(raw), item)
    days_90 = _days_90(flags, loc_warn)
    years_3 = _years_3(flags, ha, profile)
    do_not = _do_not(flags)
    questions = _questions(flags, country)
    costs = _costs(ha, p)

    if why:
        unknowns.append(why)

    return {
        "available": True,
        "suitability": suitability,
        "location_warning": loc_warn,
        "climate_grade": clim.get("grade"),
        "location_level": loc.get("level"),
        "area_ha": round(ha, 2) if ha else None,
        "risks": risks,
        "opportunities": opportunities,
        "days_90": days_90,
        "years_3": years_3,
        "do_not": do_not,
        "questions": questions,
        "cost_bands": costs,
        "unknowns": unknowns,
        "sources": sources,
        "profile_note": (f"Ecological hints use the configurable profile for {region}; "
                         "every planting or land-use idea still needs local verification."),
        "note": DATA_NOTE,
    }


def _suitability(clim: dict, flags: set, ha: float | None, loc_warn: str | None) -> str:
    grade = clim.get("grade") or "unknown"
    size = f" about {ha:.1f} ha" if ha else ""
    if grade == "unknown":
        base = f"No climate grade yet{size}. The plan below is a checklist of what to ask, not a recommendation."
    elif grade in ("excellent", "good"):
        base = (f"Climate grade {grade}{size}: the public layers do not show severe heat, water stress, "
                "fire or flood at this reading. That is not a green light — only a starting point.")
    elif grade == "caution":
        base = (f"Climate grade caution{size}: one or more risks deserve a local look before any "
                "restoration money is spent.")
    else:
        base = (f"Climate grade poor{size}: the public layers show serious constraints. "
                "Conservation here may still make sense, but only with local advice and tempered goals.")
    if {"deep_flood", "water_but_floods"} & flags:
        base += " Nearby water coincides with flood depth — riparian ideas need a hydrologist, not a planting list."
    if loc_warn:
        base += " Position is not exact: treat every local risk as provisional."
    return base


def _risks(clim: dict, flags: set, raw: dict) -> list[dict]:
    out = []
    reasons = {r for r in (clim.get("reasons") or [])}
    def add(key, title, text):
        out.append({"key": key, "title": title, "text": text})

    if "extreme_heat" in flags or "too_hot" in flags:
        bit = next((r for r in reasons if "hot" in r.lower() or "°C" in r or "35" in r),
                   "Future heat is high on the public maps.")
        add("heat", "Heat", bit + " Shade, species choice and worker safety all need a local forester.")
    if "severe_water_stress" in flags or "high_water_stress" in flags:
        bit = next((r for r in reasons if "water stress" in r.lower()), "Water stress is high by 2080.")
        add("water_stress", "Water stress", bit + " Irrigation rights and drought plans are unknown here.")
    if {"burnt_here", "repeated_burns", "fires_nearby", "high_fire_danger", "very_high_fire_danger"} & flags:
        bit = next((r for r in reasons if "fire" in r.lower() or "burnt" in r.lower()),
                   "Fire history or projected fire danger is elevated.")
        add("fire", "Fire", bit + " Fuel management and access for fire services are local questions.")
    if {"flood_zone", "deep_flood", "water_but_floods", "home_in_flood_zone"} & flags:
        depth = raw.get("flood_m")
        add("flood", "Flood",
            (f"In or near the 100-year flood layer ({depth:.1f} m)." if depth
             else "Flood exposure is flagged.") + " Do not plant or build in the floodway without advice.")
    if not out and clim.get("grade") not in (None, "unknown"):
        add("none_severe", "No severe flag from the maps",
            "The climate grade is not poor, but absence of a flag is not proof of safety.")
    if clim.get("grade") == "unknown":
        add("unknown", "Climate data missing",
            "Install or refresh the climate layers, or verify the location, then reopen this plan.")
    return out


def _opportunities(flags: set, raw: dict, country: str, profile: dict, flood,
                   item: dict | None = None) -> list[dict]:
    out = []
    wet = raw.get("water_km")
    if wet is not None and wet <= 1.0 and not flood:
        out.append({"key": "riparian", "title": "Riparian restoration (verify locally)",
                    "text": f"Permanent water about {wet:.1f} km away with no flood flag at this reading. "
                            "A local ecologist should confirm whether a buffer, fencing or native bank "
                            "vegetation is appropriate — this is not a planting prescription."})
    if wet is not None and wet <= 1.0 and flood:
        out.append({"key": "riparian_conflict", "title": "Water present, but it floods",
                    "text": "Do not treat this as a simple riparian planting opportunity. Ask about "
                            "flood regime and legal buffers first."})
    if "severe_water_stress" not in flags:
        out.append({"key": "woodland", "title": "Native woodland / natural regeneration (verify locally)",
                    "text": "Where drought stress is not extreme on the maps, natural regeneration or "
                            "native woodland may be discussed with a forester. Species choice is local."})
    if country == "PT" and profile.get("suggest_cork_oak"):
        out.append({"key": "cork_oak", "title": "Cork / oak potential — verify locally",
                    "text": "Your stewardship profile mentions cork or oak for this region. That is a "
                            "hint only: soil, exposition, cadastre use and regional plans decide. "
                            "Ask a local forester before any planting."})
    out.append({"key": "soil", "title": "Soil and habitat restoration",
                "text": "Soil cover, erosion control and habitat corridors are usually worth asking "
                        "the municipality and an ecologist about, regardless of climate grade."})
    if {"burnt_here", "repeated_burns", "fires_nearby"} & flags:
        out.append({"key": "regen", "title": "Post-fire natural regeneration",
                    "text": "After burns, natural regeneration and fuel-break design matter more than "
                            "quick replanting. Ask the fire service and a forester what is allowed."})
    text = f"{(item or {}).get('title') or ''} {(item or {}).get('description') or ''}"
    from common import has_term
    from scoring import EUCALYPTUS_WORDS
    if has_term(text, EUCALYPTUS_WORDS, negations=False):
        out.append({"key": "eucalyptus", "title": "Replace eucalyptus with natives (verify locally)",
                    "text": "The notice mentions eucalyptus. New planting of Eucalyptus spp. is normally "
                            "restricted in Portugal (DL 96/2013); conversion to native woodland is the "
                            "habitat case, not a timber crop. Ask ICNF before any felling or replanting."})
    return out


def _days_90(flags: set, loc_warn: str | None) -> list[str]:
    steps = [
        "Confirm the exact parcel on the ground and in the land registry (boundaries, access, water).",
        "Walk the land once in dry weather and once after rain; note tracks, wet patches, slopes and neighbours.",
        "Ask the municipality what land-use class and any protected or fire-management rules apply.",
    ]
    if loc_warn:
        steps.insert(0, "Verify the map position in the app (coordinates, address or cadastre) before any site work.")
    if {"burnt_here", "fires_nearby", "high_fire_danger", "very_high_fire_danger"} & flags:
        steps.append("Ask the local fire service about access routes, fuel breaks and seasonal restrictions.")
    if {"flood_zone", "deep_flood", "water_but_floods"} & flags:
        steps.append("Ask about flood history and any riparian or watercourse setbacks before fencing or planting.")
    steps.append("Do not plant, clear or dig until a local adviser has seen the place.")
    return steps


def _years_3(flags: set, ha: float | None, profile: dict) -> list[str]:
    steps = [
        "Year 1: establish access, baseline photos, and a simple management diary; no large plantings yet.",
        "Years 1–2: soil cover and erosion control where needed; monitor natural regeneration.",
        "Years 2–3: only then consider a small, locally approved planting or habitat trial on a fraction of the land.",
    ]
    if ha and ha >= 5:
        steps.append("On larger holdings, phase work by compartment so one failure does not cost the whole site.")
    if profile.get("conservation_priority"):
        steps.append(f"Your profile priority ({profile['conservation_priority']}): keep that as a goal, "
                     "but let local advice set the methods.")
    if "severe_water_stress" in flags:
        steps.append("Prefer drought-tolerant approaches and avoid irrigation-dependent schemes.")
    return steps


def _do_not(flags: set) -> list[str]:
    out = [
        "Do not treat this plan as permission to plant, clear, fence watercourses or change land use.",
        "Do not invent protected-area status, water rights or municipal classifications — ask.",
        "Do not plant non-native or fire-prone species because a generic list suggested them.",
    ]
    if {"flood_zone", "deep_flood", "water_but_floods"} & flags:
        out.append("Do not plant or build in the indicated flood area without a hydrologist or the municipality.")
    if {"burnt_here", "high_fire_danger", "very_high_fire_danger"} & flags:
        out.append("Do not replant dense fuel ladders after fire without the fire service's view.")
    if "severe_water_stress" in flags:
        out.append("Do not plan irrigation-heavy restoration where water stress is already severe on the maps.")
    return out


def _questions(flags: set, country: str) -> list[str]:
    q = [
        "What is the official land-use / soil classification, and what may be planted or left to regenerate?",
        "Are there watercourse buffers, Natura / protected designations, or hunting rights on this parcel?",
        "Who maintains the access track, and is there a registered right of way?",
    ]
    if country == "PT":
        q.append("Is there a Plano Municipal de Defesa da Floresta Contra Incêndios constraint on this land?")
    q.append("What is the flood and watercourse setback practice here, and who grants consent?")
    if {"burnt_here", "fires_nearby", "high_fire_danger", "very_high_fire_danger"} & flags:
        q.append("What fuel management and access does the fire service expect from a private owner?")
    q.append("Would a forester or ecologist visit before any spending on plants or earthworks?")
    return q


def _costs(ha: float | None, p: dict) -> list[dict]:
    """Rough €/ha planning allowances from Settings — not quotes."""
    bands = p.get("cost_eur_per_ha") or {}
    scale = ha or 1.0
    note = "Planning allowance only (Settings → Stewardship). Not a quote; local prices vary widely."
    out = []
    for key, label in (("baseline", "Baseline survey & advice"),
                       ("soil_cover", "Soil cover / erosion control"),
                       ("regen", "Natural regeneration support"),
                       ("planting_trial", "Small planting trial (if locally approved)")):
        pair = bands.get(key) or (0, 0)
        low, high = pair if isinstance(pair, (list, tuple)) and len(pair) == 2 else (0, 0)
        out.append({"key": key, "label": label,
                    "low": round(low * scale), "high": round(high * scale),
                    "eur_per_ha": (low, high), "note": note + (f" Scaled to {ha:.1f} ha." if ha else
                                                               " Per hectare until area is known.")})
    return out


def as_markdown(item: dict, plan: dict) -> str:
    if not plan.get("available"):
        return f"# Stewardship plan\n\nNot available: {plan.get('reason')}\n"
    title = (item.get("title") or item.get("id") or "listing")[:120]
    lines = [
        "# Climate & Land Stewardship Plan",
        "",
        f"**{title}** · `{item.get('id')}`",
        "",
        f"> {plan['note']}",
        "",
        "## Suitability",
        plan["suitability"],
        "",
    ]
    if plan.get("location_warning"):
        lines += ["## Location confidence", plan["location_warning"], ""]
    lines.append("## Climate risks")
    for r in plan["risks"]:
        lines += [f"### {r['title']}", r["text"], ""]
    lines.append("## Opportunities")
    for o in plan["opportunities"]:
        lines += [f"### {o['title']}", o["text"], ""]
    lines += ["## First 90 days", ""] + [f"1. {s}" for s in plan["days_90"]] + [""]
    lines += ["## First 3 years", ""] + [f"1. {s}" for s in plan["years_3"]] + [""]
    lines += ["## Do not", ""] + [f"- {s}" for s in plan["do_not"]] + [""]
    lines += ["## Questions for local advisers", ""] + [f"- {s}" for s in plan["questions"]] + [""]
    lines += ["## Rough cost categories (planning allowances)", ""]
    for c in plan["cost_bands"]:
        lines.append(f"- **{c['label']}:** €{c['low']:,}–€{c['high']:,} — {c['note']}".replace(",", " "))
    lines += ["", "## Unknowns", ""] + [f"- {u}" for u in plan["unknowns"]] + [""]
    lines += ["## Sources / dates", ""] + [f"- {s['label']}: {s['when']}" for s in plan["sources"]]
    lines += ["", plan.get("profile_note") or "", ""]
    return "\n".join(lines)


def as_text(item: dict, plan: dict) -> str:
    """Printable text without blank lines (for text_pdf)."""
    md = as_markdown(item, plan)
    return "\n".join(line if line.strip() else "·" for line in md.splitlines())


def as_docx(item: dict, plan: dict) -> bytes:
    from docx import Document
    from io import BytesIO
    doc = Document()
    doc.add_heading("Climate & Land Stewardship Plan", 0)
    doc.add_paragraph((item.get("title") or item.get("id") or "")[:200])
    doc.add_paragraph(plan.get("note") or DATA_NOTE)
    if not plan.get("available"):
        doc.add_paragraph(plan.get("reason") or "")
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()
    doc.add_heading("Suitability", level=1)
    doc.add_paragraph(plan["suitability"])
    if plan.get("location_warning"):
        doc.add_heading("Location confidence", level=1)
        doc.add_paragraph(plan["location_warning"])
    for section, key in (("Climate risks", "risks"), ("Opportunities", "opportunities")):
        doc.add_heading(section, level=1)
        for row in plan[key]:
            doc.add_heading(row["title"], level=2)
            doc.add_paragraph(row["text"])
    for section, key in (("First 90 days", "days_90"), ("First 3 years", "years_3"),
                         ("Do not", "do_not"), ("Questions for local advisers", "questions")):
        doc.add_heading(section, level=1)
        for s in plan[key]:
            doc.add_paragraph(s, style="List Bullet")
    doc.add_heading("Rough cost categories (planning allowances)", level=1)
    for c in plan["cost_bands"]:
        doc.add_paragraph(f"{c['label']}: €{c['low']:,}–€{c['high']:,} — {c['note']}".replace(",", " "))
    doc.add_heading("Unknowns", level=1)
    for u in plan["unknowns"]:
        doc.add_paragraph(u, style="List Bullet")
    doc.add_heading("Sources / dates", level=1)
    for s in plan["sources"]:
        doc.add_paragraph(f"{s['label']}: {s['when']}", style="List Bullet")
    if plan.get("profile_note"):
        doc.add_paragraph(plan["profile_note"])
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
